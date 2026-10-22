"""
Builds docs/raptor_framework.pptx -- the technical client deck for the Raptor
Deployment Gate -- directly on top of docs/reference_template.pptx so it
inherits the client's actual presentation format: the Ericsson theme colors,
"Ericsson Hilda" / "Ericsson Hilda Light" fonts, the master-level logo, and
the template's own slide layouts (Title Page, Blank, White logo end slide).

Deliberately condensed to 5 slides total (1 title + 4 dense content slides):
  1. Title
  2. Overview -- the problem, and how the gate runs (commit to cluster)
  3. The decision engine -- 9 metrics, the health-score formula, 2 overrides
  4. Three real deployments -- all three scenarios side by side
  5. Why it matters + closing contract statement

Unlike the earlier docs/raptor_framework_story.pptx (built via
build_client_ppt.py from a from-scratch theme), this version is written to be
technical-specific: every figure, formula, config value, CLI flag, and JSON
example on these slides is pulled directly from the gate/ source code and
config, not paraphrased from the HTML narrative alone. Sources:

  - gate/cli.py            -> CLI flags, modes, exit-code contract
  - gate/metrics.py        -> the 9 canonical MetricSpecs (source/unit/threshold)
  - gate/model.py          -> Isolation Forest calibration formula
  - gate/scoring.py        -> combined health-score formula, signals 1-4
  - gate/signals/trend.py  -> persistent-breach override (signal 3)
  - gate/signals/volatility.py -> window-volatility override (signal 4)
  - gate/config/production.yaml, collectors.yaml -> actual threshold/penalty values
  - deploy/k8s/deployment.yaml, service.yaml, Dockerfile -> runtime facts
  - data/live_results_*.png -> the three real captured scenario dashboards

Run:
    .venv\\Scripts\\python.exe scripts\\build_raptor_framework_pptx.py
"""
import io
import re
import zipfile
from pathlib import Path
import shutil

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
TEMPLATE = ROOT / "docs" / "reference_template.pptx"
OUT = ROOT / "docs" / "raptor_framework.pptx"
OUT_TMP = ROOT / "docs" / "raptor_framework.tmp.pptx"

# --------------------------------------------------------- Ericsson palette --
# Pulled from docs/reference_template.pptx's own theme1.xml <a:clrScheme>.
DK1 = RGBColor(0x18, 0x18, 0x18)       # theme dk1 / dk2
ACCENT = RGBColor(0x11, 0x74, 0xE6)    # theme accent1
TEAL = RGBColor(0x23, 0x96, 0x9A)      # theme accent2
VIOLET = RGBColor(0x97, 0x6C, 0xF4)    # theme accent3
GOLD = RGBColor(0xFA, 0xD2, 0x2D)      # theme accent4
ORANGE = RGBColor(0xE6, 0x6E, 0x19)    # theme accent5
ROSE = RGBColor(0xE6, 0x5D, 0x6A)      # theme accent6
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# Semantic colors (PROMOTE / ROLLBACK read universally; kept distinct from the
# theme's own accents, matching how the live dashboards color-code verdicts).
GREEN = RGBColor(0x15, 0x8A, 0x4D)
GREEN_BG = RGBColor(0xE8, 0xF7, 0xEE)
RED = ROSE
RED_BG = RGBColor(0xFC, 0xEC, 0xEE)
AMBER = ORANGE
AMBER_BG = RGBColor(0xFD, 0xF2, 0xDF)
ACCENT_BG = RGBColor(0xEA, 0xF1, 0xFD)
ACCENT_DARK = RGBColor(0x0B, 0x4E, 0xA8)
MUTED = RGBColor(0x55, 0x55, 0x55)
PANEL = RGBColor(0xF4, 0xF5, 0xF6)
BORDER = RGBColor(0xDD, 0xDD, 0xDD)
FOOT_LINE = RGBColor(0xE2, 0xE2, 0xE2)
MONO_BG = RGBColor(0x1E, 0x24, 0x2B)
MONO_FG = RGBColor(0xDC, 0xE6, 0xF2)
MONO_GREEN = RGBColor(0x7E, 0xE7, 0x9A)
MONO_ROSE = RGBColor(0xF0, 0x8C, 0x97)

FONT = "Ericsson Hilda"
FONT_LIGHT = "Ericsson Hilda Light"
MONO_FONT = "Consolas"


def _strip_template_slides(src_path: Path) -> bytes:
    """Return the template .pptx bytes with its own demo slides physically
    removed (slide parts, their rels, their Content_Types overrides, and the
    sldIdLst/presentation-rels entries pointing at them).

    We do this surgery at the zip level -- rather than via python-pptx's
    sldIdLst.remove() + drop_rel() -- because python-pptx's package keeps the
    original slide Part objects registered even after the relationship is
    dropped. Re-adding new slides afterward then reuses the freed partnames
    (slide1.xml, slide2.xml, ...) while the old orphaned parts are still
    written on save, producing a corrupt .pptx with duplicate part names.
    Deleting the slide entries before python-pptx ever opens the file avoids
    that failure mode: nothing is orphaned, nothing collides.
    """
    with zipfile.ZipFile(src_path, "r") as zin:
        pres_xml = zin.read("ppt/presentation.xml").decode("utf-8")
        rels_xml = zin.read("ppt/_rels/presentation.xml.rels").decode("utf-8")
        ct_xml = zin.read("[Content_Types].xml").decode("utf-8")

        rids = re.findall(r'<p:sldId\b[^>]*\br:id="([^"]+)"', pres_xml)
        pres_xml = re.sub(r"<p:sldIdLst>.*?</p:sldIdLst>", "<p:sldIdLst/>",
                           pres_xml, flags=re.S)

        slide_targets = []
        for rid in rids:
            m = re.search(
                rf'<Relationship\b[^>]*\bId="{re.escape(rid)}"[^>]*\bTarget="([^"]+)"[^>]*/>',
                rels_xml,
            )
            if m:
                slide_targets.append(m.group(1))
            rels_xml = re.sub(
                rf'<Relationship\b[^>]*\bId="{re.escape(rid)}"[^>]*/>', "", rels_xml
            )

        slide_parts = {f"ppt/{t}" for t in slide_targets}
        slide_rels_parts = {
            f"ppt/slides/_rels/{Path(t).name}.rels" for t in slide_targets
        }

        for part in slide_parts:
            partname = "/" + part
            ct_xml = re.sub(
                rf'<Override\b[^>]*\bPartName="{re.escape(partname)}"[^>]*/>',
                "",
                ct_xml,
            )

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                name = item.filename
                if name in slide_parts or name in slide_rels_parts:
                    continue
                if name == "ppt/presentation.xml":
                    zout.writestr(item, pres_xml)
                elif name == "ppt/_rels/presentation.xml.rels":
                    zout.writestr(item, rels_xml)
                elif name == "[Content_Types].xml":
                    zout.writestr(item, ct_xml)
                else:
                    zout.writestr(item, zin.read(name))
        return buf.getvalue()


# ---------------------------------------------------------------- load base --
# Strip the template's own demo slides (CNAaas content) before python-pptx
# ever opens it, so we inherit only the masters/layouts/theme (and the
# master-level logo graphic they carry) with zero orphaned slide parts.
_clean_template = io.BytesIO(_strip_template_slides(TEMPLATE))
prs = Presentation(_clean_template)

SLIDE_W = prs.slide_width
SLIDE_H = prs.slide_height
MARGIN_X = Inches(0.6)
CONTENT_W = SLIDE_W - 2 * MARGIN_X

LAYOUTS = {l.name: l for l in prs.slide_masters[0].slide_layouts}
BLANK = LAYOUTS["Blank"]
TITLE_PAGE = LAYOUTS["Title Page"]
END_SLIDE = LAYOUTS["White logo end slide"]

PAGE = {"n": 0}
# 5-slide deck: title (unnumbered) + 4 numbered content slides (overview,
# decision engine, three real deployments, why it matters/closing).
TOTAL_CONTENT_SLIDES = 4


def add_slide(layout=BLANK):
    return prs.slides.add_slide(layout)


def set_bg(slide, color=None):
    if color is None:
        return
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def _no_autosize(tf):
    try:
        tf.word_wrap = True
        tf.auto_size = None
    except Exception:
        pass


def letter_spaced(run, pts_hundredths=120):
    rPr = run._r.get_or_add_rPr()
    rPr.set("spc", str(pts_hundredths))


def textbox(slide, left, top, width, height, text, size=16, bold=False,
            color=DK1, align=PP_ALIGN.LEFT, font=FONT, italic=False,
            anchor=MSO_ANCHOR.TOP, line_spacing=1.15, spaced=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    _no_autosize(tf)
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        r.font.name = font
        if spaced:
            letter_spaced(r)
    return box


def bullet_list(slide, left, top, width, height, items, size=13.5,
                 color=DK1, dot_color=ACCENT, line_spacing=1.25, gap_pt=8,
                 font=FONT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    _no_autosize(tf)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_after = Pt(gap_pt)
        r1 = p.add_run()
        r1.text = "\u25CF  "
        r1.font.size = Pt(size)
        r1.font.color.rgb = dot_color
        r1.font.name = font
        r2 = p.add_run()
        r2.text = item
        r2.font.size = Pt(size)
        r2.font.color.rgb = color
        r2.font.name = font
    return box


def rect(slide, left, top, width, height, fill=WHITE, line=BORDER, line_w=1.0,
          shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06):
    shp = slide.shapes.add_shape(shape, left, top, width, height)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def card(slide, left, top, width, height, title=None, body=None,
         fill=WHITE, line=BORDER, title_color=DK1, title_size=14,
         body_color=MUTED, body_size=11.5, align=PP_ALIGN.LEFT, pad=0.2,
         valign=MSO_ANCHOR.TOP, title_bold=True, line_w=1.0, font=FONT,
         mono_body=False):
    shp = rect(slide, left, top, width, height, fill=fill, line=line, line_w=line_w)
    tf = shp.text_frame
    _no_autosize(tf)
    tf.margin_left = Inches(pad)
    tf.margin_right = Inches(pad)
    tf.margin_top = Inches(pad * 0.8)
    tf.margin_bottom = Inches(pad * 0.8)
    tf.vertical_anchor = valign
    first = True
    if title:
        p0 = tf.paragraphs[0]
        p0.alignment = align
        r = p0.add_run()
        r.text = title
        r.font.size = Pt(title_size)
        r.font.bold = title_bold
        r.font.color.rgb = title_color
        r.font.name = font
        first = False
    if body:
        for line in body.split("\n"):
            p = tf.add_paragraph() if not first else tf.paragraphs[0]
            first = False
            p.alignment = align
            p.line_spacing = 1.18
            p.space_before = Pt(4) if title else Pt(0)
            r = p.add_run()
            r.text = line
            r.font.size = Pt(body_size)
            r.font.color.rgb = body_color
            r.font.name = MONO_FONT if mono_body else font
    return shp


def stat_chip(slide, left, top, width, height, label, value, color=ACCENT,
              fill=WHITE, line=BORDER):
    shp = rect(slide, left, top, width, height, fill=fill, line=line)
    tf = shp.text_frame
    _no_autosize(tf)
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.16)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    r0.text = label.upper()
    r0.font.size = Pt(10)
    r0.font.bold = True
    r0.font.color.rgb = MUTED
    r0.font.name = FONT
    letter_spaced(r0, 50)
    p1 = tf.add_paragraph()
    p1.space_before = Pt(2)
    r1 = p1.add_run()
    r1.text = value
    r1.font.size = Pt(15.5)
    r1.font.bold = True
    r1.font.color.rgb = color
    r1.font.name = FONT
    return shp


def code_block(slide, left, top, width, height, lines, size=11.5):
    """A terminal/JSON style block: dark bg, monospace, optional color runs.

    `lines` is a list of either plain strings, or (text, color) tuples so
    JSON keys/values and ROLLBACK/PROMOTE can be highlighted.
    """
    shp = rect(slide, left, top, width, height, fill=MONO_BG, line=None,
               shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
    tf = shp.text_frame
    _no_autosize(tf)
    tf.margin_left = Inches(0.22)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.16)
    tf.margin_bottom = Inches(0.16)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.18
        segs = line if isinstance(line, list) else [(line, MONO_FG)]
        for text, color in segs:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.name = MONO_FONT
            r.font.color.rgb = color
    return shp


def h_arrow(slide, left, top, width=Inches(0.4), height=Inches(0.2), color=MUTED):
    shp = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def v_arrow(slide, left, top, width=Inches(0.2), height=Inches(0.28), color=MUTED):
    shp = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, left, top, width, height)
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def number_badge(slide, cx, cy, diameter, text, fill=ACCENT, color=WHITE, size=14):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx, cy, diameter, diameter)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = color
    r.font.name = FONT
    return shp


def picture_framed(slide, path, left, top, height, line=BORDER):
    pic = slide.shapes.add_picture(str(path), left, top, height=height)
    rect(slide, pic.left, pic.top, pic.width, pic.height, fill=None,
         line=line, line_w=1.0, shape=MSO_SHAPE.RECTANGLE)
    return pic


def header(slide, kicker, title, accent=ACCENT):
    textbox(slide, MARGIN_X, Inches(0.35), Inches(10), Inches(0.28),
            kicker.upper(), size=11, bold=True, color=accent, spaced=True)
    textbox(slide, MARGIN_X, Inches(0.64), CONTENT_W, Inches(0.55),
            title, size=25, bold=True, color=DK1, font=FONT_LIGHT)
    rect(slide, MARGIN_X, Inches(1.28), CONTENT_W, Pt(1.4), fill=BORDER,
         line=None, shape=MSO_SHAPE.RECTANGLE)


def footer(slide, brand="RAPTOR DEPLOYMENT GATE  \u00b7  TECHNICAL BRIEFING"):
    PAGE["n"] += 1
    rect(slide, MARGIN_X, Inches(7.05), CONTENT_W, Pt(1), fill=FOOT_LINE, line=None,
         shape=MSO_SHAPE.RECTANGLE)
    textbox(slide, MARGIN_X, Inches(7.13), Inches(9), Inches(0.28), brand,
            size=9, color=MUTED, spaced=True)
    textbox(slide, SLIDE_W - MARGIN_X - Inches(1.5), Inches(7.13), Inches(1.5), Inches(0.28),
            f"{PAGE['n']:02d} / {TOTAL_CONTENT_SLIDES:02d}", size=9, color=MUTED,
            align=PP_ALIGN.RIGHT)


def clean_table(slide, left, top, width, height, rows_data, col_widths,
                 header_fill=ACCENT, header_color=WHITE, body_size=12,
                 header_size=12, row_colors=None, align_cols=None, font=FONT):
    n_rows = len(rows_data)
    n_cols = len(rows_data[0])
    gfx = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    tbl = gfx.table
    tbl.first_row = False
    tbl.horz_banding = False
    for ci, w in enumerate(col_widths):
        tbl.columns[ci].width = w
    for ri, row in enumerate(rows_data):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.text = str(val)
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(header_size if ri == 0 else body_size)
            p.font.name = font
            p.font.bold = (ri == 0)
            p.alignment = (align_cols[ci] if align_cols else PP_ALIGN.LEFT)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.12)
            cell.margin_right = Inches(0.12)
            cell.margin_top = Inches(0.05)
            cell.margin_bottom = Inches(0.05)
            cell.fill.solid()
            if ri == 0:
                cell.fill.fore_color.rgb = header_fill
                p.font.color.rgb = header_color
            else:
                bg = PANEL if (row_colors is None and ri % 2 == 0) else WHITE
                cell.fill.fore_color.rgb = row_colors[ri] if row_colors else bg
                p.font.color.rgb = DK1
    return tbl


def fill_placeholder(slide, idx, text):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            ph.text_frame.text = text
            return ph
    return None


# =============================================================================
# SLIDE 1 -- Title (uses the template's own "Title Page" layout + logo)
# =============================================================================
s = add_slide(TITLE_PAGE)
fill_placeholder(s, 0, "Raptor Deployment Gate")
fill_placeholder(
    s, 1,
    "Technical briefing: architecture, scoring engine, and live results\n"
    "eric-cenx-rest-api \u00b7 robot-poc namespace \u00b7 October 2026",
)
fill_placeholder(s, 10, "Engineering Briefing")
fill_placeholder(s, 11, "Raptor Deployment Gate")
fill_placeholder(s, 12, "2026-10-22")

# =============================================================================
# SLIDE 2 -- The problem & how it runs (merges former problem + architecture
# slides into one dense overview slide)
# =============================================================================
s = add_slide(); set_bg(s)
header(s, "01 \u00b7 Overview", "Why this gate exists, and how it runs")

col_w = Inches(5.5)
gap = Inches(0.3)
left_x = MARGIN_X
right_x = MARGIN_X + col_w + gap

rect(s, left_x, Inches(1.45), col_w, Inches(0.42), fill=RED_BG, line=RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, left_x + Inches(0.2), Inches(1.54), col_w - Inches(0.4), Inches(0.26),
        "Today, without automation", size=12, bold=True, color=RED)
bullet_list(s, left_x, Inches(1.98), col_w, Inches(1.8), [
    "A person watches dashboards after every deploy and decides",
    "Judgment varies by who's on call, and subtle issues get missed",
    "A recurring minor issue can keep slipping through, release after release",
], dot_color=RED, size=11.5, gap_pt=5)

rect(s, right_x, Inches(1.45), col_w, Inches(0.42), fill=GREEN_BG, line=GREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, right_x + Inches(0.2), Inches(1.54), col_w - Inches(0.4), Inches(0.26),
        "With the Raptor Gate", size=12, bold=True, color=GREEN)
bullet_list(s, right_x, Inches(1.98), col_w, Inches(1.8), [
    "The same 9 metrics, scored the same way, on every deployment",
    "Isolation Forest + SLA engine + two cross-sample override signals",
    "One contract: process exit code 0 (PROMOTE) or 1 (ROLLBACK)",
], dot_color=GREEN, size=11.5, gap_pt=5)

# -- compact architecture strip: commit to cluster --
textbox(s, MARGIN_X, Inches(3.75), CONTENT_W, Inches(0.28),
        "FROM COMMIT TO CLUSTER \u2014 deploy/k8s/deployment.yaml, applied to namespace robot-poc",
        size=10, bold=True, color=ACCENT, spaced=True)

steps = [
    ("CI/CD + Kaniko", "GitLab CI / Jenkins triggers an\nunprivileged in-cluster image build"),
    ("Deployment/raptor-gate", "2 replicas \u00b7 ClusterIP :80\u2192:8080\nIsolation Forest baked into the image"),
    ("Collectors", "Prometheus, Kubernetes API, Redis,\nlive HTTP probe \u2014 9 metrics every pass"),
    ("Service: eric-cenx-rest-api", "the canary workload\nbeing evaluated in robot-poc"),
]
n = len(steps)
box_w = Inches(2.68)
arrow_w = Inches(0.35)
total_w = box_w * n + arrow_w * (n - 1)
start_x = MARGIN_X + (CONTENT_W - total_w) / 2
y = Inches(4.15)
x = start_x
for i, (title, sub) in enumerate(steps):
    card(s, x, y, box_w, Inches(1.25), title=title, body=sub, fill=ACCENT_BG, line=ACCENT,
         title_color=ACCENT_DARK, title_size=11.5, body_color=MUTED, body_size=9.5,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    x += box_w
    if i < n - 1:
        h_arrow(s, x + Inches(0.015), y + Inches(0.53), width=arrow_w - Inches(0.07))
        x += arrow_w

card(s, MARGIN_X, Inches(5.65), CONTENT_W, Inches(1.0),
     body="Resources: requests 100m CPU / 256Mi \u00b7 limits 500m CPU / 512Mi. "
          "Probes: readiness/liveness GET /health (5s/10s and 10s/20s). "
          "Collection failure \u21d2 exit code 1 immediately \u2014 the gate never scores a partial sample.",
     fill=PANEL, line=BORDER, body_color=MUTED, body_size=11, valign=MSO_ANCHOR.MIDDLE,
     align=PP_ALIGN.CENTER)
footer(s)

# =============================================================================
# SLIDE 3 -- The decision engine (merges metrics table + formula + overrides)
# =============================================================================
s = add_slide(); set_bg(s)
header(s, "02 \u00b7 The Decision Engine", "Nine metrics, one formula, two overrides")

metrics_rows = [
    ("Metric", "SLA", "Metric", "SLA"),
    ("cpu_usage", "\u2264 80%", "availability", "\u2265 99.5%"),
    ("memory_usage", "\u2264 75%", "redis_latency", "\u2264 10 ms"),
    ("pod_restarts", "\u2264 3", "endpoint_latency", "\u2264 200 ms"),
    ("http_error_rate", "\u2264 1%", "packet_loss", "\u2264 1%"),
    ("p99_latency", "\u2264 200 ms", "", ""),
]
mt_w = Inches(5.5)
clean_table(s, MARGIN_X, Inches(1.5), mt_w, Inches(2.5), metrics_rows,
            col_widths=[Inches(1.9), Inches(0.85), Inches(1.9), Inches(0.85)],
            header_size=10.5, body_size=10.5,
            align_cols=[PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.CENTER])
textbox(s, MARGIN_X, Inches(4.05), mt_w, Inches(0.5),
        "gate/metrics.py \u2014 single source of truth for the SLA engine, "
        "model feature order, and CSV loaders. Threshold = 70, penalty = 20 "
        "points/breach (gate/config/production.yaml).",
        size=9.5, color=MUTED, line_spacing=1.2)

form_x = MARGIN_X + mt_w + Inches(0.35)
form_w = CONTENT_W - mt_w - Inches(0.35)
code_block(s, form_x, Inches(1.5), form_w, Inches(1.9), [
    [("anomaly_health = 100 * (1 - deficit/tolerance)", MONO_FG)],
    [("deductions = 20 * len(violations)", MONO_FG)],
    [("health = clamp(anomaly - deductions, 0, 100)", MONO_FG)],
    [("", MONO_FG)],
    [("health < 70  ", MONO_ROSE), ("\u2192 ROLLBACK", MONO_ROSE)],
    [("health \u2265 70  ", MONO_GREEN), ("\u2192 provisional PROMOTE", MONO_GREEN)],
], size=10.5)

card(s, form_x, Inches(3.55), form_w, Inches(1.0),
     title="3 \u00b7 Persistent breach", body="Same metric breaches \u2265 2 of last 5 "
     "samples \u2192 forced ROLLBACK, regardless of health score.",
     fill=AMBER_BG, line=AMBER, title_color=AMBER, title_size=11, body_color=DK1,
     body_size=9.5, valign=MSO_ANCHOR.TOP)
card(s, form_x, Inches(4.65), form_w, Inches(1.0),
     title="4 \u00b7 Window volatility", body="\u2265 50% of raw sub-window ticks breach "
     "\u2192 flagged, even if Mean aggregation hid it from the SLA engine.",
     fill=AMBER_BG, line=AMBER, title_color=AMBER, title_size=11, body_color=DK1,
     body_size=9.5, valign=MSO_ANCHOR.TOP)

card(s, MARGIN_X, Inches(5.85), CONTENT_W, Inches(0.8),
     body="Both overrides are additive and opt-in on top of the base formula \u2014 "
          "Verdict.forced_rollback marks when 3 or 4, not the 70-point "
          "threshold, drove the decision.",
     fill=PANEL, line=BORDER, body_color=MUTED, body_size=10.5, valign=MSO_ANCHOR.MIDDLE,
     align=PP_ALIGN.CENTER)
footer(s)

# =============================================================================
# SLIDE 4 -- Three real deployments (merges divider + 3 scenario slides into
# one 3-column proof slide)
# =============================================================================
s = add_slide(); set_bg(s)
header(s, "03 \u00b7 Proof, Not a Demo Script", "Three real deployments, three outcomes")

textbox(s, MARGIN_X, Inches(1.45), CONTENT_W, Inches(0.3),
        "Scored live by gate.cli against eric-cenx-rest-api in robot-poc on October 21, 2026 \u2014 not simulated.",
        size=10.5, color=MUTED)

scenarios = [
    dict(
        letter="A", headline="Clean release", verdict="PROMOTE", color=GREEN, bg=GREEN_BG,
        img=DATA / "live_results_eric-cenx-rest-api_20261021_110134_promote.png",
        stats=[("Samples", "82"), ("Verdict", "82/82 PROMOTE"), ("Health range", "78.4\u201393.5"), ("Breaches", "0")],
        note="No metric breaches twice in the 5-sample window \u2014 neither override has anything to catch.",
    ),
    dict(
        letter="B", headline="Drift + flicker", verdict="MOSTLY PROMOTE", color=AMBER, bg=AMBER_BG,
        img=DATA / "live_results_eric-cenx-rest-api_20261021_110731_change_values.png",
        stats=[("Samples", "63"), ("Verdict", "60 PROMOTE \u00b7 3 ROLLBACK"), ("Health range", "51.9\u201383.9"), ("Breaches", "3 \u00d7 redis")],
        note="Redis spikes cost a 20-pt deduction each time, but never repeat \u2192 signal 3 never fires.",
    ),
    dict(
        letter="C", headline="Degraded canary", verdict="ROLLBACK", color=RED, bg=RED_BG,
        img=DATA / "live_results_eric-cenx-rest-api_20261021_130754_rollback.png",
        stats=[("Samples", "56"), ("Verdict", "41 PROMOTE \u00b7 15 ROLLBACK"), ("Health range", "51.9\u201378.5"), ("Breaches", "12 \u00d7 p99 \u00b7 3 \u00d7 redis")],
        note="p99_latency recurs well past 2 of 5 samples \u2014 the persistent-breach override forces several ROLLBACKs outright.",
    ),
]

col_w = Inches(3.95)
col_gap = Inches(0.22)
total_w = col_w * 3 + col_gap * 2
start_x = MARGIN_X + (CONTENT_W - total_w) / 2
x = start_x
for sc in scenarios:
    rect(s, x, Inches(1.85), col_w, Inches(0.42), fill=sc["bg"], line=sc["color"], line_w=1.0)
    number_badge(s, x + Inches(0.08), Inches(1.9), Inches(0.32), sc["letter"], fill=sc["color"], color=WHITE, size=11)
    textbox(s, x + Inches(0.48), Inches(1.9), col_w - Inches(1.55), Inches(0.32), sc["headline"], size=11, bold=True, color=DK1)
    textbox(s, x + col_w - Inches(1.65), Inches(1.9), Inches(1.6), Inches(0.32), sc["verdict"], size=10, bold=True, color=sc["color"], align=PP_ALIGN.RIGHT)

    picture_framed(s, sc["img"], x, Inches(2.4), Inches(2.15))

    sy = Inches(4.65)
    for label, value in sc["stats"]:
        textbox(s, x, sy, Inches(1.5), Inches(0.24), label.upper(), size=7.5, bold=True, color=MUTED, spaced=True)
        textbox(s, x + Inches(1.5), sy, col_w - Inches(1.5), Inches(0.24), value, size=9.5, bold=True, color=sc["color"], align=PP_ALIGN.RIGHT)
        sy += Inches(0.24)

    card(s, x, sy + Inches(0.08), col_w, Inches(1.25), body=sc["note"], fill=PANEL, line=BORDER,
         body_color=MUTED, body_size=8.8, valign=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT)
    x += col_w + col_gap

footer(s, "Captured live against eric-cenx-rest-api, robot-poc  \u00b7  Oct 21, 2026")

# =============================================================================
# SLIDE 5 -- Why it matters + closing (merges comparison table + closing)
# =============================================================================
s = add_slide(); set_bg(s)
header(s, "04 \u00b7 Why It Matters", "The case for an automatic gate")

compare_rows = [
    ("Without an automated gate", "With the Raptor Gate"),
    ("A human watches dashboards and decides, under time pressure", "Automatic, consistent, takes seconds \u2014 exit code 0/1"),
    ("Judgment varies by who's on call", "Every deployment scored by the same model and the same 9 thresholds"),
    ("Subtle cross-metric issues are easy to miss", "The Isolation Forest is built to catch combinations no single rule would flag"),
    ("A single-value Mean can hide a majority-bad window", "Signal 4 checks the raw sub-window ticks directly"),
    ("A minor, recurring issue can keep slipping through", "Signal 3 forces a rollback once a metric repeats \u2265 2 of the last 5 samples"),
]
half_col = Emu(int(CONTENT_W / 2))
tbl = clean_table(s, MARGIN_X, Inches(1.5), CONTENT_W, Inches(3.9), compare_rows,
                   col_widths=[half_col, half_col], header_size=12, body_size=10.5,
                   header_fill=DK1)
for ri in range(len(compare_rows)):
    for ci in range(2):
        cell = tbl.cell(ri, ci)
        if ri == 0:
            cell.fill.fore_color.rgb = RED if ci == 0 else GREEN
            cell.text_frame.paragraphs[0].font.color.rgb = WHITE
        else:
            cell.fill.fore_color.rgb = RED_BG if ci == 0 else GREEN_BG
            cell.text_frame.paragraphs[0].font.color.rgb = DK1

card(s, MARGIN_X, Inches(5.65), CONTENT_W, Inches(1.0),
     title="One contract: exit code 0 to promote, exit code 1 to roll back.",
     body="Any pipeline \u2014 GitLab CI, Jenkins, or a custom script \u2014 can branch on that directly. No proprietary platform required.",
     fill=ACCENT_BG, line=ACCENT, title_color=ACCENT_DARK, title_size=13, body_color=DK1,
     body_size=10.5, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
footer(s)

# ------------------------------------------------------------------------- save
prs.save(str(OUT_TMP))
try:
    shutil.move(str(OUT_TMP), str(OUT))
    print(f"Saved: {OUT}  ({len(prs.slides._sldIdLst)} slides)")
except PermissionError:
    print(f"Target locked; wrote to {OUT_TMP} instead. Close any app with the pptx open and retry, "
          f"or rename {OUT_TMP.name} to {OUT.name} manually.")
